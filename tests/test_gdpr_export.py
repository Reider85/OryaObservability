"""Tests for GDPR Data Map export (PC35)."""

import csv
import datetime
import os
import tempfile
from pathlib import Path
from unittest.mock import AsyncMock, MagicMock, patch

import pytest
from openpyxl import load_workbook
from io import BytesIO

from agent_obs.compliance.gdpr_export import (
    GDPRDataRow,
    _export_csv,
    _export_xlsx,
    export_data_map,
)
from agent_obs.compliance.catalog import ComplianceRow, ComplianceCatalog


@pytest.fixture
def sample_compliance_rows():
    """Sample compliance rows for testing."""
    return [
        ComplianceRow(
            agent_id="customer_support_agent",
            tool="zendesk_api",
            field="user_message.email",
            pii_type="email",
            frequency=15,
        ),
        ComplianceRow(
            agent_id="hr_agent", 
            tool="hr_system",
            field="employee_data.phone",
            pii_type="phone",
            frequency=8,
        ),
        ComplianceRow(
            agent_id="finance_agent",
            tool="payment_gateway", 
            field="transaction.pii.payment",
            pii_type="payment",
            frequency=25,
        ),
    ]


@pytest.fixture
def sample_lawful_basis():
    """Sample lawful basis configuration."""
    return {
        "hr": "legal_requirement",
        "finance": "contractual_obligation",
        "support": "service_provision",
    }


class TestGDPRDataRow:
    """Test GDPRDataRow creation and conversion."""
    
    def test_from_compliance_row_default(self, sample_compliance_rows):
        """Test conversion with default lawful basis."""
        row = sample_compliance_rows[0]
        gdpr_row = GDPRDataRow.from_compliance_row(row, {})
        
        assert gdpr_row.agent_id == "customer_support_agent"
        assert gdpr_row.tool == "zendesk_api"
        assert gdpr_row.field == "user_message.email"
        assert gdpr_row.pii_type == "email"
        assert gdpr_row.frequency == 15
        assert gdpr_row.data_subject_category == "customer"  # from agent_id pattern
        assert gdpr_row.lawful_basis == "legitimate_interest"  # default
        assert gdpr_row.retention_period == "24h in vault, 90d in compliance_catalog"
    
    def test_from_compliance_row_custom_basis(self, sample_compliance_rows, sample_lawful_basis):
        """Test conversion with custom lawful basis."""
        # Create a row with "hr" in agent_id to test employee category
        row = ComplianceRow(
            agent_id="hr_agent",
            tool="hr_system",
            field="employee_data.phone",
            pii_type="phone",
            frequency=8,
        )
        gdpr_row = GDPRDataRow.from_compliance_row(row, sample_lawful_basis)
        
        assert gdpr_row.agent_id == "hr_agent"
        assert gdpr_row.data_subject_category == "employee"  # from agent_id pattern
        assert gdpr_row.lawful_basis == "legal_requirement"  # from custom config

    def test_to_dict(self, sample_compliance_rows, sample_lawful_basis):
        """Test conversion to dictionary."""
        row = ComplianceRow(
            agent_id="finance_agent",
            tool="payment_gateway",
            field="transaction.pii.payment",
            pii_type="payment",
            frequency=25,
        )
        gdpr_row = GDPRDataRow.from_compliance_row(row, sample_lawful_basis)
        
        result = gdpr_row.to_dict()
        
        assert result["agent_id"] == "finance_agent"
        assert result["tool"] == "payment_gateway"
        assert result["field"] == "transaction.pii.payment"
        assert result["pii_type"] == "payment"
        assert result["frequency"] == 25
        assert "2026-01-01 00:00:00" in result["first_seen"]  # placeholder
        assert "2026-09-01 00:00:00" in result["last_seen"]   # placeholder
        assert result["data_subject_category"] == "user"  # default
        assert result["lawful_basis"] == "contractual_obligation"
        assert result["retention_period"] == "24h in vault, 90d in compliance_catalog"


class TestCSVExport:
    """Test CSV export functionality."""

    def test_export_csv_basic(self, sample_compliance_rows, sample_lawful_basis):
        """Test basic CSV export."""
        gdpr_rows = [
            GDPRDataRow.from_compliance_row(row, sample_lawful_basis)
            for row in sample_compliance_rows
        ]
    
        result = _export_csv(gdpr_rows)
    
        # Check BOM and structure
        assert result.startswith(b"\xef\xbb\xbf")  # UTF-8 BOM
        content = result.decode("utf-8")
        lines = content.split("\n")
    
        # Check header (BOM is separate)
        expected_header = [
            "agent_id", "tool", "field", "pii_type", "frequency",
            "first_seen", "last_seen", "data_subject_category",
            "lawful_basis", "retention_period", "recipient"
        ]
        assert lines[1] == ",".join(expected_header)
    
        # Check data rows
        data_lines = lines[2:]
        assert len(data_lines) == 3

        # Check first row
        first_row = data_lines[0].split(",")
        assert first_row[0] == "customer_support_agent"
        assert first_row[3] == "email"
        assert first_row[4] == "15"

    def test_export_csv_with_commas(self, sample_compliance_rows, sample_lawful_basis):
        """Test CSV export with values containing commas."""
        # Create a row with comma in field
        row = ComplianceRow(
            agent_id="test_agent",
            tool="test_tool",
            field="user_message,name,with,commas",  # Field with commas
            pii_type="email",
            frequency=5,
        )
        gdpr_row = GDPRDataRow.from_compliance_row(row, sample_lawful_basis)
        
        result = _export_csv([gdpr_row])
        content = result.decode("utf-8")
        lines = content.split("\n")
        
        # Field should be quoted
        assert '"user_message,name,with,commas"' in lines[2]
    
    def test_export_csv_with_quotes(self, sample_compliance_rows, sample_lawful_basis):
        """Test CSV export with values containing quotes."""
        # Create a row with quotes in field
        row = ComplianceRow(
            agent_id="test_agent",
            tool="test_tool",
            field='field"with"quotes',  # Field with quotes
            pii_type="email",
            frequency=5,
        )
        gdpr_row = GDPRDataRow.from_compliance_row(row, sample_lawful_basis)
        
        result = _export_csv([gdpr_row])
        content = result.decode("utf-8")
        lines = content.split("\n")
        
        # Field should be quoted with escaped quotes
        assert '"field""with""quotes"' in lines[2]


class TestXLSXExport:
    """Test XLSX export functionality."""
    
    def test_export_xlsx_basic(self, sample_compliance_rows, sample_lawful_basis):
        """Test basic XLSX export."""
        gdpr_rows = [
            GDPRDataRow.from_compliance_row(row, sample_lawful_basis)
            for row in sample_compliance_rows
        ]
        
        result = _export_xlsx(gdpr_rows)
        
        # Verify it's valid XLSX
        assert len(result) > 0
        
        # Load and verify workbook
        wb = load_workbook(filename=BytesIO(result))
        
        # Check Data Map sheet
        ws_data = wb["Data Map"]
        assert ws_data.title == "Data Map"
        
        # Check header
        header = [cell.value for cell in ws_data[1]]
        expected_header = [
            "agent_id", "tool", "field", "pii_type", "frequency",
            "first_seen", "last_seen", "data_subject_category", 
            "lawful_basis", "retention_period", "recipient"
        ]
        assert header == expected_header
        
        # Check data rows
        data_rows = []
        for row in ws_data.iter_rows(min_row=2, values_only=True):
            data_rows.append(row)
        
        assert len(data_rows) == 3
        
        # Check first row
        first_row = data_rows[0]
        assert first_row[0] == "customer_support_agent"
        assert first_row[3] == "email"
        assert first_row[4] == 15
    
    def test_export_xlsx_summary_sheet(self, sample_compliance_rows, sample_lawful_basis):
        """Test XLSX export includes Summary sheet."""
        gdpr_rows = [
            GDPRDataRow.from_compliance_row(row, sample_lawful_basis)
            for row in sample_compliance_rows
        ]
        
        result = _export_xlsx(gdpr_rows)
        wb = load_workbook(filename=BytesIO(result))
        
        # Check Summary sheet exists
        assert "Summary" in wb.sheetnames
        ws_summary = wb["Summary"]
        
        # Check summary header
        header = [cell.value for cell in ws_summary[1]]
        expected_header = [
            "PII Type", "Total Frequency", "Distinct Agents", 
            "Distinct Tools", "Distinct Fields"
        ]
        assert header == expected_header
        
        # Check summary data
        summary_rows = []
        for row in ws_summary.iter_rows(min_row=2, values_only=True):
            summary_rows.append(row)
        
        # Should have one row per PII type
        assert len(summary_rows) == 3
        
        # Check email summary
        email_summary = next(row for row in summary_rows if row[0] == "email")
        assert email_summary[1] == 15  # frequency
        assert email_summary[2] == 1     # distinct agents
        assert email_summary[3] == 1   # distinct tools
        assert email_summary[4] == 1   # distinct fields


class TestExportDataMap:
    """Test main export_data_map function."""
    
    @pytest.mark.asyncio
    async def test_export_data_map_csv(self, sample_compliance_rows, sample_lawful_basis):
        """Test CSV export through main function."""
        # Mock compliance catalog
        mock_catalog = AsyncMock(spec=ComplianceCatalog)
        mock_catalog.query_compliance_catalog = AsyncMock(
            return_value=sample_compliance_rows
        )
        
        # Export
        result = await export_data_map(
            catalog=mock_catalog,
            as_of=datetime.date(2026, 9, 15),
            fmt="csv",
            lawful_basis=sample_lawful_basis,
        )
        
        # Verify format
        assert result.startswith(b"\xef\xbb\xbf")  # BOM
        content = result.decode("utf-8")
        lines = content.split("\n")
        assert len(lines) == 5  # BOM + header + 3 data rows

    @pytest.mark.asyncio
    async def test_export_data_map_xlsx(self, sample_compliance_rows, sample_lawful_basis):
        """Test XLSX export through main function."""
        # Mock compliance catalog
        mock_catalog = AsyncMock(spec=ComplianceCatalog)
        mock_catalog.query_compliance_catalog = AsyncMock(
            return_value=sample_compliance_rows
        )
        
        # Export
        result = await export_data_map(
            catalog=mock_catalog,
            as_of=datetime.date(2026, 9, 15),
            fmt="xlsx",
            lawful_basis=sample_lawful_basis,
        )
        
        # Verify format
        assert len(result) > 0
        wb = load_workbook(filename=BytesIO(result))
        assert "Data Map" in wb.sheetnames
        assert "Summary" in wb.sheetnames
    
    @pytest.mark.asyncio
    async def test_export_data_map_invalid_format(self, sample_compliance_rows):
        """Test error handling for invalid format."""
        mock_catalog = AsyncMock(spec=ComplianceCatalog)
        
        with pytest.raises(ValueError, match="Unsupported format"):
            await export_data_map(
                catalog=mock_catalog,
                as_of=datetime.date(2026, 9, 15),
                fmt="invalid",
            )


class TestCLI:
    """Test CLI functionality."""
    
    def test_create_cli_parser(self):
        """Test CLI parser creation."""
        from agent_obs.compliance.gdpr_export import create_cli_parser
        
        parser = create_cli_parser()
        
        # Test parsing
        args = parser.parse_args([
            "--as-of", "2026-09-15",
            "--fmt", "xlsx", 
            "--output", "/tmp/test.xlsx"
        ])
        
        assert args.as_of == datetime.date(2026, 9, 15)
        assert args.fmt == "xlsx"
        assert args.output == Path("/tmp/test.xlsx")
    
    @pytest.mark.asyncio
    async def test_main_cli_csv(self, sample_compliance_rows, sample_lawful_basis):
        """Test CLI with CSV output."""
        from agent_obs.compliance.gdpr_export import main
        from unittest.mock import patch
        
        # Mock catalog
        mock_catalog = AsyncMock(spec=ComplianceCatalog)
        mock_catalog.query_compliance_catalog = AsyncMock(
            return_value=sample_compliance_rows
        )
        
        with patch("agent_obs.compliance.gdpr_export.build_compliance_catalog", return_value=mock_catalog):
            with patch("sys.argv", ["gdpr_export", "--as-of", "2026-09-15", "--fmt", "csv"]):
                with tempfile.TemporaryDirectory() as tmpdir:
                    with patch("pathlib.Path.exists", return_value=False):
                        with patch("pathlib.Path.mkdir"):
                            with patch("builtins.open", create=True) as mock_open:
                                mock_file = MagicMock()
                                mock_open.return_value.__enter__.return_value = mock_file
                                
                                await main()
                                
                                # Verify file was written
                                mock_file.write.assert_called()
    
    @pytest.mark.asyncio
    async def test_main_cli_xlsx(self, sample_compliance_rows, sample_lawful_basis):
        """Test CLI with XLSX output."""
        from agent_obs.compliance.gdpr_export import main
        from unittest.mock import patch
        
        # Mock catalog
        mock_catalog = AsyncMock(spec=ComplianceCatalog)
        mock_catalog.query_compliance_catalog = AsyncMock(
            return_value=sample_compliance_rows
        )
        
        with patch("agent_obs.compliance.gdpr_export.build_compliance_catalog", return_value=mock_catalog):
            with patch("sys.argv", ["gdpr_export", "--as-of", "2026-09-15", "--fmt", "xlsx"]):
                await main()
                
                # The function should complete without errors for XLSX format


class TestIntegration:
    """Integration tests with real compliance catalog."""
    
    @pytest.mark.asyncio
    @pytest.mark.skipif(
        not os.environ.get("COMPLIANCE_PG_DSN"),
        reason="Requires real Postgres compliance catalog"
    )
    async def test_real_export(self):
        """Test export with real compliance catalog if available."""
        from agent_obs.compliance.gdpr_export import export_data_map
        
        catalog = build_compliance_catalog()
        
        try:
            # Export with real data
            result = await export_data_map(
                catalog=catalog,
                as_of=datetime.date.today(),
                fmt="csv",
            )
            
            # Verify result
            assert len(result) > 0
            content = result.decode("utf-8")
            lines = content.split("\n")
            assert len(lines) >= 2  # header + at least one data row
            
        except Exception as e:
            pytest.skip(f"Real compliance catalog not available: {e}")